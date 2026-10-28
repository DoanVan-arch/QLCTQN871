# -*- coding: utf-8 -*-
"""Tien ich nhap/xuat du lieu nhan su dang file Excel (.xlsx)."""
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (ma truong, nhan cot, loai: text|date|bool)
EXCEL_FIELDS = [
    ("full_name", "Họ và tên", "text"),
    ("gender", "Giới tính", "text"),
    ("ethnicity", "Dân tộc", "text"),
    ("religion", "Tôn giáo", "text"),
    ("date_of_birth", "Ngày sinh", "date"),
    ("enlistment_date", "Ngày nhập ngũ", "date"),
    ("rank", "Cấp bậc", "text"),
    ("rank_date", "TG nhận cấp bậc", "date"),
    ("position", "Chức vụ", "text"),
    ("position_date", "TG nhận chức vụ", "date"),
    ("youth_union_date", "TG vào Đoàn", "date"),
    ("youth_union_cell", "Chi đoàn", "text"),
    ("party_sympathy_class", "Đã học cảm tình Đảng", "bool"),
    ("party_join_date", "TG vào Đảng (dự bị)", "date"),
    ("party_cell", "Chi bộ", "text"),
    ("party_official_date", "TG chính thức", "date"),
    ("unit_before_enrollment", "Đơn vị trước khi nhập học", "text"),
    ("current_residence", "Nơi ở", "text"),
    ("registered_residence", "Trú quán", "text"),
    ("hometown", "Quê quán", "text"),
    ("school_name", "Tên trường theo học", "text"),
    ("major", "Chuyên ngành", "text"),
    ("education_level", "Bậc học", "text"),
    ("enrollment_date", "TG nhập học", "date"),
    ("reception_date", "TG tiếp nhận", "date"),
    ("graduation_date", "TG tốt nghiệp", "date"),
    ("rewards", "Khen thưởng", "text"),
    ("discipline", "Kỷ luật", "text"),
    ("citizen_id", "Số CCCD", "text"),
    ("officer_id", "CMSQ", "text"),
    ("phone", "Số điện thoại", "text"),
    ("contact_person", "Người báo tin", "text"),
    ("notes", "Ghi chú", "text"),
]

HEADER_FILL = PatternFill(start_color="223322", end_color="223322", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _cell_value(person, field, kind):
    value = getattr(person, field, None)
    if kind == "bool":
        return "Có" if value else "Không"
    return value


def export_personnel_to_excel(items):
    """Xuat danh sach Personnel ra Workbook openpyxl. Tra ve Workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "DanhSach"

    headers = [label for _, label, _ in EXCEL_FIELDS]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for person in items:
        row = []
        for field, _, kind in EXCEL_FIELDS:
            value = _cell_value(person, field, kind)
            if kind == "date" and isinstance(value, date):
                row.append(value.strftime("%d/%m/%Y"))
            else:
                row.append(value or "")
        ws.append(row)

    for col_idx, (_, label, _) in enumerate(EXCEL_FIELDS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, min(28, len(label) + 4))

    ws.freeze_panes = "A2"
    return wb


def build_import_template():
    """Tao file mau (chi co tieu de) de nguoi dung dien du lieu nhap."""
    return export_personnel_to_excel([])


def _parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_personnel_excel(file_stream):
    """Doc file Excel va tra ve danh sach dict du lieu nhan su (chua tao Personnel).

    Tra ve (rows, errors) voi rows la list[dict] theo ten truong cua model,
    errors la list[str] mo ta loi theo dong.
    """
    wb = load_workbook(file_stream, data_only=True)
    ws = wb.active

    rows = []
    errors = []
    header_cells = next(ws.iter_rows(min_row=1, max_row=1), [])
    header_labels = [str(c.value).strip() if c.value else "" for c in header_cells]
    label_to_field = {label: (field, kind) for field, label, kind in EXCEL_FIELDS}

    col_map = {}
    for idx, label in enumerate(header_labels):
        if label in label_to_field:
            col_map[idx] = label_to_field[label]

    for row_num, row in enumerate(ws.iter_rows(min_row=2), start=2):
        values = [c.value for c in row]
        if all(v in (None, "") for v in values):
            continue
        data = {}
        for idx, (field, kind) in col_map.items():
            if idx >= len(values):
                continue
            raw = values[idx]
            if kind == "date":
                data[field] = _parse_date(raw)
            elif kind == "bool":
                data[field] = str(raw).strip() in ("Có", "co", "Co", "True", "true", "1", "x", "X")
            else:
                data[field] = str(raw).strip() if raw is not None else ""

        if not data.get("full_name"):
            errors.append(f"Dòng {row_num}: thiếu Họ và tên, đã bỏ qua.")
            continue

        rows.append(data)

    return rows, errors
